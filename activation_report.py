"""Экспорт результатов: TXT/DOCX/CSV."""

from __future__ import annotations

from pathlib import Path
import csv
import json

import numpy as np
import pandas as pd


def export_tables_csv(results: dict, out_dir: str) -> list[str]:
    Path(out_dir).mkdir(parents=True, exist_ok=True)
    files = []
    f = results["fields"]

    table_map = {
        "T_solid_K.csv": f["Ts_2d"],
        "T_gas_K.csv": f["Tg_2d"],
        "X.csv": f["X_2d"],
        "S_BET_m2g.csv": f["S_BET_2d"],
        "P_H2O_Pa.csv": f["P_h2o"],
        "eta.csv": f["eta_2d"],
        "r_act.csv": f["r_act_2d"],
        "Q_conv.csv": f["Q_conv_2d"],
        "rho_bed.csv": f["rho_bed_2d"],
        "D_eff.csv": f["D_eff_2d"],
    }

    for name, arr in table_map.items():
        path = Path(out_dir) / name
        np.savetxt(path, arr, delimiter=",", fmt="%.6e")
        files.append(str(path))

    return files


def export_summary_csv(rows: list[dict], out_path: str) -> str:
    if not rows:
        return out_path
    keys = list(rows[0].keys())
    with open(out_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=keys)
        w.writeheader()
        for row in rows:
            w.writerow(row)
    return out_path


def generate_txt_report(results: dict, out_path: str) -> str:
    p = results["passport"]
    m = results["mass_balance"]
    e = results["energy_balance"]
    i = results["engineering_indicators"]
    status = results.get("status", "Требуется проверка параметров")
    interp = results.get("interpretation", "Интерпретация отсутствует.")

    def fmt(v, nd=3):
        try:
            return f"{float(v):.{nd}f}"
        except Exception:  # noqa: BLE001
            return "None"

    lines = [
        "=" * 72,
        "ОТЧЕТ: ПАРОВАЯ АКТИВАЦИЯ УГЛЯ (v3.2 wrapper)",
        "=" * 72,
        f"Статус режима: {results['status']}",
        "",
        "ИНТЕГРАЛЬНЫЕ РЕЗУЛЬТАТЫ",
        f"- Средний обгар X: {p['Burnoff_Mean_Pct']:.2f} %",
        f"- Выход активированного угля: {p['Yield_Pct']:.2f} %",
        f"- S_BET: {p['BET_Surface_m2g']:.1f} м2/г",
        f"- Макс. диаметр пор: {p['Max_Pore_Diameter_nm']:.2f} нм",
        f"- Средний Ts на выходе: {p['Ts_Exit_Mean_K'] - 273.15:.1f} °C",
        f"- Мин. eta: {p['Min_Eta']:.3f}",
        "",
        "БАЛАНСЫ",
        f"- Невязка массы (solid): {m['solid_error_pct']:.4f} %",
        f"- Невязка массы (gas): {m['gas_error_pct']:.4f} %",
        f"- Энергия с паром: {e['Q_gas_W']/1000:.2f} кВт",
        f"- Энергия с сырьем: {e['Q_solid_W']/1000:.2f} кВт",
        f"- Сток тепла на эндотермику: {abs(e['Q_rxn_W'])/1000:.2f} кВт",
        "",
        "ИНЖЕНЕРНЫЕ ИНДЕКСЫ (требуют экспериментальной калибровки)",
        f"- I_act = S_BET / X: {fmt(i.get('I_act'))}",
        f"- K_ret = Y_AU * S_BET / 100: {fmt(i.get('K_ret'))}",
        f"- strength_factor (model): {fmt(i.get('strength_factor_model', i.get('strength_factor')))}",
        f"- K_quality = S_BET * strength_factor / X: {fmt(i.get('K_quality'))}",
        f"- Примечание: {i['label']}",
        "",
        "ПРЕДУПРЕЖДЕНИЯ",
    ]
    lines.extend([f"- {w}" for w in results.get("warnings", [])] or ["- Нет"])
    lines.extend([
        "",
        "ФИЗИЧЕСКАЯ ВАЛИДАЦИЯ РАСЧЁТА",
        f"- Статус режима: {status}",
        "- Инженерные индексы:",
        f"  * I_act: {i.get('I_act')}",
        f"  * K_ret: {i.get('K_ret')}",
        f"  * K_quality: {i.get('K_quality')}",
        f"  * strength_factor: {i.get('strength_factor_model', i.get('strength_factor'))}",
        f"  * K_T: {i.get('K_T')}",
        f"- Интерпретация: {interp}",
    ])
    parametric = results.get("parametric", {})
    if parametric:
        lines.extend(["", "ПАРАМЕТРИЧЕСКИЙ АНАЛИЗ"])
        for key, data in parametric.items():
            if isinstance(data, list) and data:
                lines.append(f"- {key}: выполнено {len(data)} точек.")
                lines.append(f"  Тренд (первый/последний S_BET): {data[0].get('S_BET', 'n/a')} -> {data[-1].get('S_BET', 'n/a')}")
            else:
                lines.append(f"- {key}: данных нет.")

    opt_rows = results.get("optimization", [])
    constraints = results.get("optimization_constraints", {})
    if opt_rows:
        best = opt_rows[0]
        lines.extend([
            "",
            "ПОДБОР РАЦИОНАЛЬНОГО РЕЖИМА",
            f"- Ограничения: {constraints}",
            f"- Лучший режим: T_steam={best.get('T_steam')}, steam_flow={best.get('steam_flow')}, rho={best.get('rho_granule')}, d={best.get('d_particle')}",
            f"- X={best.get('X_mean')}, Y={best.get('Y_AU')}, S_BET={best.get('S_BET')}, gas_residual={best.get('gas_residual')}",
            f"- Статус: {best.get('status')}, warnings_count={best.get('warnings_count')}, score={best.get('score')}",
        ])
    lines.append("=" * 72)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))

    return out_path


def generate_docx_report(results: dict, out_path: str) -> str | None:
    try:
        from docx import Document
    except Exception:
        return None

    doc = Document()
    doc.add_heading("Отчет по паровой активации", level=1)
    doc.add_paragraph(f"Статус режима: {results['status']}")

    p = results["passport"]
    for k, v in p.items():
        doc.add_paragraph(f"{k}: {v}")

    doc.add_heading("Предупреждения", level=2)
    for w in results.get("warnings", []):
        doc.add_paragraph(w, style="List Bullet")

    doc.add_heading("Физическая валидация расчёта", level=2)
    doc.add_paragraph(f"Статус режима: {results.get('status', 'Требуется проверка параметров')}")
    idx = results.get("engineering_indicators", {})
    doc.add_paragraph("Инженерные индексы:")
    doc.add_paragraph(f"I_act: {idx.get('I_act')}", style="List Bullet")
    doc.add_paragraph(f"K_ret: {idx.get('K_ret')}", style="List Bullet")
    doc.add_paragraph(f"K_quality: {idx.get('K_quality')}", style="List Bullet")
    doc.add_paragraph(f"strength_factor: {idx.get('strength_factor_model', idx.get('strength_factor'))}", style="List Bullet")
    doc.add_paragraph(f"K_T: {idx.get('K_T')}", style="List Bullet")
    doc.add_paragraph(f"Интерпретация: {results.get('interpretation', 'Интерпретация отсутствует.')}")

    parametric = results.get("parametric", {})
    if parametric:
        doc.add_heading("Параметрический анализ", level=2)
        for key, data in parametric.items():
            if isinstance(data, list):
                doc.add_paragraph(f"{key}: {len(data)} точек", style="List Bullet")

    opt_rows = results.get("optimization", [])
    if opt_rows:
        best = opt_rows[0]
        doc.add_heading("Подбор рационального режима", level=2)
        doc.add_paragraph(f"Ограничения: {results.get('optimization_constraints', {})}")
        doc.add_paragraph(f"Лучший режим: {best}")

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


def save_results_json(results_jsonable: dict, out_path: str) -> str:
    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(results_jsonable, f, ensure_ascii=False, indent=2)
    return out_path


def generate_figure_caption(figure_key: str, context: dict) -> str:
    captions = {
        "feedstock_comparison": "Рисунок 4.5.1 — Сравнение расчётных показателей активации для семи видов сырья при едином базовом режиме.",
        "burnoff_by_feedstock": "Рисунок 4.5.2 — Средняя степень обгара карбонизатов различных видов сырья.",
        "activation_index": "Рисунок 4.5.3 — Индекс активационной восприимчивости I_act для исследуемых видов сырья.",
        "temperature_sweep": "Рисунок 4.5.4 — Влияние температуры пара на интегральные характеристики активации.",
        "steam_flow_sweep": "Рисунок 4.5.5 — Влияние расхода пара на степень обгара, выход и развитость поверхности.",
        "particle_size_sweep": "Рисунок 4.5.6 — Влияние размера частиц на диффузионные ограничения и степень активации.",
        "density_sweep": "Рисунок 4.5.7 — Влияние плотности гранул на показатели активации и условную прочность.",
        "field_Ts": "Рисунок 4.5.8 — Распределение температуры твёрдой фазы Ts(h,l).",
        "field_Tg": "Рисунок 4.5.9 — Распределение температуры газовой фазы Tg(h,l).",
        "field_X": "Рисунок 4.5.10 — Пространственное поле степени обгара X(h,l).",
        "field_S_BET": "Рисунок 4.5.11 — Пространственное поле удельной поверхности S_BET(h,l).",
        "field_P_H2O": "Рисунок 4.5.12 — Распределение парциального давления H2O в слое.",
        "field_eta": "Рисунок 4.5.13 — Поле фактора эффективности η(h,l).",
        "field_r_act": "Рисунок 4.5.14 — Поле скорости реакции активации r_act(h,l).",
        "field_rho_bed": "Рисунок 4.5.15 — Поле насыпной плотности слоя ρ_bed(h,l).",
        "field_D_eff": "Рисунок 4.5.16 — Поле эффективной диффузии D_eff(h,l).",
        "optimization_result": "Рисунок 4.5.17 — Ранжирование режимов по интегральному критерию score.",
    }
    return captions.get(figure_key, f"Рисунок — {figure_key}")


def generate_result_interpretation(result_type: str, data, params: dict) -> str:
    base = (
        "Расчёт показывает тенденции в рамках текущей модели и заданных параметров. "
        "Полученные значения являются модельной оценкой и требуют экспериментального уточнения."
    )
    if isinstance(data, pd.DataFrame) and not data.empty:
        if result_type == "temperature_sweep":
            return (
                f"{base} При варьировании температуры пара наблюдается изменение X_mean и S_BET "
                "в соответствии с усилением активации при росте температуры."
            )
        if result_type == "feedstock_compare":
            return (
                f"{base} Сравнение сырья показывает различия по I_act и K_ret, "
                "что может быть использовано для предварительного выбора сырьевой базы."
            )
        if result_type == "optimization":
            best = data.iloc[0].to_dict()
            return (
                f"{base} По результатам перебора лучший режим имеет score={best.get('score')} "
                f"и параметры T_steam={best.get('T_steam')}, steam_flow={best.get('steam_flow')}."
            )
    return base


def get_dissertation_graph_templates() -> list[dict]:
    return [
        {"id": 1, "title": "Выход АУ и S_BET для 7 видов сырья", "before_text": "На рисунке представлена расчётно-экспериментальная зависимость, полученная по результатам серии сравнения сырья.", "after_text": "График показывает относительную перспективность сырья; при замене предварительных данных на экспериментальные текст должен быть уточнён."},
        {"id": 2, "title": "Средняя степень обгара для 7 видов сырья", "before_text": "Результаты получены расчётным сравнением при едином базовом режиме.", "after_text": "Показана чувствительность обгара к типу сырья; вывод требует калибровки по эксперименту."},
        {"id": 3, "title": "Индекс активационной восприимчивости", "before_text": "Индекс рассчитан как I_act=S_BET/X по модельным данным.", "after_text": "Индекс отражает удельное развитие поверхности на единицу обгара и служит инженерным критерием предварительного выбора."},
        {"id": 4, "title": "Влияние температуры пара на выход и йодную активность", "before_text": "Представлена расчётно-экспериментальная зависимость по температурной серии.", "after_text": "Тенденция отражает усиление активации с ростом температуры; численные оценки требуют экспериментального уточнения, включая йодную активность."},
        {"id": 5, "title": "Влияние размера частиц на выход и йодную активность", "before_text": "Данные получены модельной серией по размеру частиц.", "after_text": "График демонстрирует влияние диффузионных ограничений; оценка йодной активности является предварительной."},
        {"id": 6, "title": "Влияние плотности гранул и обгара на активность по толуолу", "before_text": "Построение выполнено на основе расчётных данных плотностной серии.", "after_text": "Показано, что рост плотности может ограничивать доступность пор; интерпретация по толуолу требует экспериментальной подстановки."},
        {"id": 7, "title": "Влияние степени обгара на прочность гранул", "before_text": "Использован условный strength_factor в качестве инженерного индикатора.", "after_text": "Вывод о прочности является оценочным и подлежит обязательной экспериментальной верификации."},
        {"id": 8, "title": "Связь степени обгара и S_BET", "before_text": "График построен по выходным сериям модели.", "after_text": "Наблюдается характерная область рационального развития поверхности; параметры плато требуют уточнения."},
        {"id": 9, "title": "Температура угля, реактора и пара", "before_text": "Профили температур получены из расчёта полей реактора.", "after_text": "Показана тепловая согласованность режима при заданных условиях."},
        {"id": 10, "title": "Влияние температуры пиролиза на последующую активацию", "before_text": "Использован сценарный анализ входных параметров карбонизата.", "after_text": "Расчётная оценка указывает на значимость исходного состояния карбонизата; вывод требует экспериментального уточнения."},
        {"id": 11, "title": "Влияние плотности прессования на последующую активацию", "before_text": "Построено по серии изменения плотности гранул.", "after_text": "Показаны предварительные зависимости между плотностью, обгаром и поверхностью."},
        {"id": 12, "title": "Пространственная неоднородность слоя и газовой фазы", "before_text": "На рисунке представлены двумерные расчётные поля реактора.", "after_text": "График демонстрирует неоднородность параметров по высоте и толщине слоя; при замене предварительных данных на экспериментальные текст должен быть уточнён."},
    ]


def _criteria_table(results: dict) -> pd.DataFrame:
    p = results.get("passport", {})
    m = results.get("mass_balance", {})
    return pd.DataFrame(
        [
            ["Выход активированного угля", p.get("Yield_Pct"), "%", "из расчёта"],
            ["Степень обгара", p.get("Burnoff_Mean_Pct"), "%", "из расчёта"],
            ["S_BET", p.get("BET_Surface_m2g"), "м²/г", "из расчёта"],
            ["Йодная активность", "требует калибровки", "мг/г", "placeholder"],
            ["Активность по толуолу", "требует калибровки", "%", "placeholder"],
            ["Прочность", "требует калибровки", "%", "placeholder"],
            ["Газовая невязка", m.get("gas_error_pct"), "%", "из расчёта"],
            ["Тепловой баланс (Q_rxn)", results.get("energy_balance", {}).get("Q_rxn_W"), "Вт", "из расчёта"],
        ],
        columns=["Критерий", "Значение", "Ед.", "Источник"],
    )


def _rational_modes_table(feedstock_df: pd.DataFrame | None, optimization_df: pd.DataFrame | None) -> pd.DataFrame:
    rows = []
    if feedstock_df is not None and not feedstock_df.empty:
        for _, r in feedstock_df.iterrows():
            rows.append({
                "feedstock": r.get("feedstock"),
                "recommended_T_steam": "требует калибровки",
                "recommended_steam_flow": "требует калибровки",
                "recommended_X_range": "30–40%",
                "expected_S_BET": r.get("S_BET"),
                "limiting_factor": r.get("top_warning", "нет"),
                "recommended_application": "сорбенты общего назначения (предв. оценка)",
            })
    if not rows and optimization_df is not None and not optimization_df.empty:
        for _, r in optimization_df.head(5).iterrows():
            rows.append({
                "feedstock": "базовый сценарий",
                "recommended_T_steam": r.get("T_steam"),
                "recommended_steam_flow": r.get("steam_flow"),
                "recommended_X_range": "30–40%",
                "expected_S_BET": r.get("S_BET"),
                "limiting_factor": r.get("top_warning", "нет"),
                "recommended_application": "требует калибровки",
            })
    return pd.DataFrame(rows)


def generate_dissertation_report(results, params, parametric_results=None, optimization_results=None, output_path=None):
    output_path = output_path or str(Path(params.get("meta", {}).get("output_dir", "section_2_5_output")) / "dissertation_4_5_report.docx")
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)

    feedstock_df = None if parametric_results is None else parametric_results.get("feedstock_compare")
    optimization_df = optimization_results

    criteria_df = _criteria_table(results)
    feedstock_cmp_df = feedstock_df if isinstance(feedstock_df, pd.DataFrame) else pd.DataFrame()
    rational_df = _rational_modes_table(feedstock_cmp_df, optimization_df)
    opt_top_df = optimization_df.head(10) if isinstance(optimization_df, pd.DataFrame) and not optimization_df.empty else pd.DataFrame()

    payload = dict(results)
    payload["parametric"] = {}
    if isinstance(parametric_results, dict):
        payload["parametric"] = {k: (v.to_dict(orient="records") if isinstance(v, pd.DataFrame) else v) for k, v in parametric_results.items()}
    payload["optimization"] = [] if optimization_df is None else optimization_df.to_dict(orient="records")
    payload["optimization_constraints"] = params.get("targets", {})

    # TXT always
    txt_path = str(out.with_suffix(".txt"))
    generate_txt_report(payload, txt_path)

    # CSV tables
    criteria_df.to_csv(out.parent / "4_5_table_criteria.csv", index=False, encoding="utf-8-sig")
    if not feedstock_cmp_df.empty:
        feedstock_cmp_df.to_csv(out.parent / "4_5_table_feedstock_compare.csv", index=False, encoding="utf-8-sig")
    if not rational_df.empty:
        rational_df.to_csv(out.parent / "4_5_table_rational_modes.csv", index=False, encoding="utf-8-sig")
    if not opt_top_df.empty:
        opt_top_df.to_csv(out.parent / "4_5_table_optimization_top10.csv", index=False, encoding="utf-8-sig")

    docx_path = None
    if out.suffix.lower() == ".docx":
        if generate_docx_report(payload, str(out)):
            docx_path = str(out)

    return {
        "txt": txt_path,
        "docx": docx_path,
        "criteria_csv": str(out.parent / "4_5_table_criteria.csv"),
        "feedstock_csv": None if feedstock_cmp_df.empty else str(out.parent / "4_5_table_feedstock_compare.csv"),
        "rational_csv": None if rational_df.empty else str(out.parent / "4_5_table_rational_modes.csv"),
        "optimization_csv": None if opt_top_df.empty else str(out.parent / "4_5_table_optimization_top10.csv"),
    }
